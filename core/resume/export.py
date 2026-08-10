"""Resume export: plaintext-ATS, DOCX, PDF.

Renders a ParsedResume to a clean, ATS-friendly layout. DOCX via python-docx,
PDF via PyMuPDF (both already dependencies) — no new libraries.
"""
from __future__ import annotations

from pathlib import Path


def resume_to_text(parsed) -> str:
    lines: list[str] = []
    c = parsed.contact or {}
    if c.get("name"):
        lines.append(c["name"])
    contact_line = " | ".join(x for x in (c.get("email"), c.get("phone"), c.get("location")) if x)
    if contact_line:
        lines.append(contact_line)
    for link in parsed.links or []:
        lines.append(str(link))
    if parsed.summary:
        lines += ["", "SUMMARY", parsed.summary]
    if parsed.experience:
        lines += ["", "EXPERIENCE"]
        for e in parsed.experience:
            head = " — ".join(x for x in (e.title, e.company) if x)
            dates = " – ".join(x for x in (e.start, e.end) if x)
            lines.append(f"{head}" + (f" ({dates})" if dates else "")
                         + (f", {e.location}" if e.location else ""))
            for b in e.bullets or []:
                lines.append(f"  • {b}")
    if parsed.education:
        lines += ["", "EDUCATION"]
        for e in parsed.education:
            lines.append(" — ".join(x for x in (e.degree, e.institution) if x)
                         + (f" ({e.year})" if e.year else ""))
    tech = parsed.all_skills()
    if tech:
        lines += ["", "SKILLS", ", ".join(tech)]
    if parsed.certifications:
        lines += ["", "CERTIFICATIONS", ", ".join(str(x) for x in parsed.certifications)]
    return "\n".join(lines).strip() + "\n"


def export(parsed, output_path: str, fmt: str = "txt") -> str:
    fmt = (fmt or "txt").lower()
    text = resume_to_text(parsed)
    path = Path(output_path).with_suffix(f".{fmt}")
    if fmt == "docx":
        _export_docx(parsed, path)
    elif fmt == "pdf":
        _export_pdf(text, path)
    else:
        path.write_text(text, encoding="utf-8")
    return str(path.resolve())


def _export_docx(parsed, path: Path) -> None:
    from docx import Document
    doc = Document()
    c = parsed.contact or {}
    if c.get("name"):
        doc.add_heading(c["name"], level=0)
    contact_line = " | ".join(x for x in (c.get("email"), c.get("phone"), c.get("location")) if x)
    if contact_line:
        doc.add_paragraph(contact_line)
    if parsed.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(parsed.summary)
    if parsed.experience:
        doc.add_heading("Experience", level=1)
        for e in parsed.experience:
            head = " — ".join(x for x in (e.title, e.company) if x)
            dates = " – ".join(x for x in (e.start, e.end) if x)
            # bold is a run/font property — setting it on the Paragraph is a no-op.
            run = doc.add_paragraph().add_run(head + (f" ({dates})" if dates else ""))
            run.bold = True
            for b in e.bullets or []:
                doc.add_paragraph(str(b), style="List Bullet")
    if parsed.education:
        doc.add_heading("Education", level=1)
        for e in parsed.education:
            doc.add_paragraph(" — ".join(x for x in (e.degree, e.institution) if x)
                              + (f" ({e.year})" if e.year else ""))
    tech = parsed.all_skills()
    if tech:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(tech))
    if parsed.certifications:
        doc.add_heading("Certifications", level=1)
        doc.add_paragraph(", ".join(str(x) for x in parsed.certifications))
    doc.save(str(path))


def _export_pdf(text: str, path: Path) -> None:
    import fitz  # PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(56, 56, page.rect.width - 56, page.rect.height - 56)
    # insert_textbox flows text and returns overflow; add pages as needed.
    remaining = text
    while remaining:
        leftover = page.insert_textbox(rect, remaining, fontsize=10.5,
                                       fontname="helv", align=0)
        if leftover >= 0:  # non-negative => it all fit (returns leftover height)
            break
        # negative means overflow; split roughly and continue on a new page
        half = len(remaining) // 2
        page.insert_textbox(rect, remaining[:half], fontsize=10.5, fontname="helv")
        remaining = remaining[half:]
        page = doc.new_page()
    doc.save(str(path))
    doc.close()
