"""Phase 4 — resume module tests.

LLM-dependent paths (structured extraction, tailoring) are tested through their
deterministic branches: the local fallback, the safety net, and the rubric math.
"""
from __future__ import annotations

import pytest

import db
from core.resume import (
    ParsedResume,
    check_ats,
    export,
    expand_term,
    expansion_map,
    gap_report,
    resume_to_text,
    score_against_job,
    seniority_band,
    target_titles,
    weighted_skills,
    years_experience,
)
from core.resume.generate import _safe_rewrite, tailor_bullets
from core.resume.schema import Experience


@pytest.fixture
def dbpath(tmp_path, monkeypatch):
    from db import connection
    p = tmp_path / "test.db"
    monkeypatch.setattr(connection, "DB_PATH", p)
    monkeypatch.setattr(db, "DB_PATH", p, raising=False)
    db.init_db()
    return p


def _resume():
    return ParsedResume(
        contact={"name": "Jane Doe", "email": "jane@example.com", "location": "San Francisco, CA"},
        summary="Senior backend engineer with 7 years building REST APIs.",
        experience=[Experience(title="Senior Software Engineer", company="Acme",
                               start="2019", end="Present",
                               bullets=["Built Python microservices on AWS.",
                                        "Led PostgreSQL migration."])],
        education=[__import__("core.resume.schema", fromlist=["Education"]).Education(
            institution="UC Berkeley", degree="B.S. Computer Science", year="2018")],
        skills={"technical": ["Python", "Django", "AWS", "PostgreSQL"],
                "soft": [], "tools": ["Docker"], "languages": []},
        certifications=["AWS Certified Solutions Architect"])


# --------------------------------------------------------------------------- #
class TestSchema:
    def test_roundtrip(self):
        r = _resume()
        r2 = ParsedResume.from_dict(r.to_dict())
        assert r2.contact["email"] == "jane@example.com"
        assert r2.experience[0].title == "Senior Software Engineer"
        assert r2.all_skills()[:2] == ["Python", "Django"]

    def test_flattened_skills_tolerated(self):
        r = ParsedResume.from_dict({"skills": ["Python", "Go"]})
        assert r.skills["technical"] == ["Python", "Go"]


class TestKeywords:
    def test_synonym_expansion(self):
        assert set(expand_term("k8s")) == {"k8s", "kubernetes"}
        assert expansion_map(["k8s", "python", "unmatched"]) == {
            "k8s": ["k8s", "kubernetes"], "python": ["python", "py"]}

    def test_weighted_skills_order(self):
        ws = dict(weighted_skills(_resume()))
        assert ws["Python"] == 1.0 and ws["Docker"] == 0.7
        assert ws["AWS Certified Solutions Architect"] == 0.9

    def test_years_from_summary_fallback(self):
        assert years_experience(_resume()) == 7.0     # dates + summary "7 years"

    def test_seniority_and_targets(self):
        assert seniority_band(_resume()) == "senior"
        assert "Senior Software Engineer" in target_titles(_resume())


class TestATS:
    def test_checks_and_score(self):
        text = resume_to_text(_resume()) + " experience education skills " * 60
        res = check_ats(_resume(), text, "jane_doe.pdf")
        names = {c["name"]: c["ok"] for c in res["checks"]}
        assert names["Contact info present"] is True
        assert names["Parseable structure"] is True
        assert 0 <= res["score"] <= 100


class TestRubric:
    def test_explainable_dimensions(self):
        job = {"title": "Senior Python Engineer",
               "description": "5+ years Python, Django, AWS, PostgreSQL. Bachelor's required. REST.",
               "work_mode": "remote", "seniority": "senior"}
        r = score_against_job(_resume(), job, ats_result={"score": 80})
        assert 0 <= r["composite"] <= 100
        assert set(r["dimensions"]) >= {"hard_skills", "seniority", "education", "ats"}
        assert r["dimensions"]["seniority"]["score"] == 100      # senior vs senior
        assert r["dimensions"]["hard_skills"]["score"] > 50
        assert "Python" in r["strengths"] or "python" in [s.lower() for s in r["strengths"]]

    def test_missing_skills_surfaced(self):
        job = {"title": "ML Engineer", "description": "TensorFlow PyTorch CUDA required."}
        r = score_against_job(_resume(), job)
        assert any("tensorflow" in m.lower() or "pytorch" in m.lower() for m in r["missing"])

    def test_absent_dimensions_are_neutral_not_guessed(self):
        job = {"title": "Engineer", "description": "Join our team."}
        r = score_against_job(_resume(), job)
        # no explicit skills/years/degree -> those dims are None, excluded from composite
        assert r["dimensions"]["experience_years"]["score"] is None


class TestGenerateSafety:
    def test_reverts_invented_metrics(self):
        orig = "Built microservices on AWS."
        bad = "Built microservices on AWS handling 5M requests, cutting latency 40%."
        final, flagged, reason = _safe_rewrite(orig, bad)
        assert flagged and final == orig and "40" in reason

    def test_accepts_honest_rephrase(self):
        orig = "Built Python microservices on AWS."
        good = "Developed Python microservices leveraging AWS infrastructure."
        final, flagged, _ = _safe_rewrite(orig, good, {"python", "aws"})
        assert not flagged and final == good

    def test_reverts_invented_skill(self):
        # Rewrite adds "Kubernetes" the candidate never listed -> reverted + flagged.
        orig = "Built Python microservices on AWS."
        bad = "Built Python microservices on AWS and Kubernetes."
        final, flagged, reason = _safe_rewrite(orig, bad, allowed_skills={"python", "aws"})
        assert flagged and final == orig and "kubernetes" in reason.lower()

    def test_surfacing_owned_skill_is_allowed(self):
        # Candidate HAS Docker (in résumé) -> surfacing it into a bullet is honest.
        orig = "Built Python microservices."
        good = "Built Python microservices, containerized with Docker."
        final, flagged, _ = _safe_rewrite(orig, good, allowed_skills={"python", "docker"})
        assert not flagged and final == good

    def test_tailor_noop_without_llm(self):
        diffs = tailor_bullets(_resume(), {"title": "X", "description": "y"}, use_llm=False)
        assert len(diffs) == 2 and all(not d.changed for d in diffs)

    def test_gap_report(self):
        g = gap_report({"missing": ["Kubernetes"], "dimensions": {"hard_skills": {"score": 30}}})
        assert g["missing"] == ["Kubernetes"] and "hard_skills" in g["weak_dimensions"]


class TestExport:
    def test_all_formats(self, tmp_path):
        r = _resume()
        assert "Jane Doe" in resume_to_text(r)
        for fmt in ("txt", "docx", "pdf"):
            path = export(r, str(tmp_path / "resume"), fmt)
            assert path.endswith(fmt)
            import os
            assert os.path.getsize(path) > 0


class TestReviewFixes:
    """Regression tests for defects found by the adversarial review workflow."""

    def test_ats_sections_word_boundary(self):
        # 'framework'/'teamwork' must NOT satisfy the 'work' section heading.
        text = "framework teamwork network coursework " * 30
        res = check_ats(ParsedResume(), text, "r.pdf")
        names = {c["name"]: c["ok"] for c in res["checks"]}
        assert names["Standard section headings"] is False

    def test_hard_skills_count_based_not_weight_capped(self):
        # Candidate has both required skills, but as low-weight tools -> still 100%.
        r = ParsedResume(skills={"technical": [], "tools": ["Docker", "Kubernetes"],
                                 "soft": [], "languages": []})
        res = score_against_job(r, {"title": "Engineer", "description": "Docker and Kubernetes required."})
        assert res["dimensions"]["hard_skills"]["score"] == 100

    def test_years_requirement_needs_context(self):
        r = _resume()
        # "2 years ago" is not a requirement -> years dimension stays neutral.
        res = score_against_job(r, {"title": "Eng", "description": "We were founded 2 years ago."})
        assert res["dimensions"]["experience_years"]["score"] is None
        res2 = score_against_job(r, {"title": "Eng", "description": "Requires 5+ years."})
        assert res2["dimensions"]["experience_years"]["score"] is not None

    def test_set_default_bad_id_keeps_a_default(self, dbpath):
        from core.resume import save_resume, set_default, default_resume
        save_resume("A", "x " * 50, _resume())
        set_default(9999)  # non-existent
        assert default_resume() is not None and default_resume()["is_default"]

    def test_delete_default_promotes_survivor(self, dbpath):
        from core.resume import save_resume, set_default, default_resume
        from core.resume.store import delete_resume
        a = save_resume("A", "x " * 50, _resume())
        b = save_resume("B", "y " * 50, _resume())
        set_default(a)
        delete_resume(a)                       # deleting the default
        d = default_resume()
        assert d is not None and d["id"] == b  # survivor promoted

    def test_from_dict_survives_null_arrays(self):
        # A stray JSON null must not crash the parse or wipe good data.
        d = {"contact": {"email": "a@b.com"}, "skills": {"technical": ["Python"], "tools": None},
             "experience": None, "projects": None, "education": [{"degree": None}]}
        r = ParsedResume.from_dict(d)
        assert r.contact["email"] == "a@b.com"
        assert r.all_skills() == ["Python"]          # null 'tools' didn't crash all_skills
        assert r.experience == []

    def test_from_dict_skills_as_string(self):
        r = ParsedResume.from_dict({"skills": "Python, Go, Rust"})
        assert r.skills["technical"] == ["Python", "Go", "Rust"]

    def test_dotted_skill_invention_caught(self):
        orig = "Built web apps."
        bad = "Built web apps with ASP.NET."   # .net not in résumé
        final, flagged, reason = _safe_rewrite(orig, bad, allowed_skills={"python"})
        assert flagged and final == orig

    def test_spelled_number_fabrication_caught(self):
        orig = "Scaled the platform."
        bad = "Scaled the platform to five million users."
        final, flagged, _ = _safe_rewrite(orig, bad)
        assert flagged and final == orig

    def test_education_partial_credit_for_non_degree(self):
        from core.resume.schema import Education
        r = ParsedResume(education=[Education(degree="Coding Bootcamp Certificate")])
        res = score_against_job(r, {"title": "Eng", "description": "Bachelor's degree required."})
        assert res["dimensions"]["education"]["score"] == 70   # not full, not zero

    def test_location_none_does_not_crash(self):
        r = ParsedResume(contact={"location": None})
        res = score_against_job(r, {"title": "Eng", "description": "onsite role", "work_mode": "onsite",
                                    "location_country": "USA"})
        assert res["dimensions"]["location"]["score"] in (None, 40, 70, 100)

    def test_docx_export_bold_is_on_run(self, tmp_path):
        # The title run must actually be bold (was silently set on the paragraph).
        from docx import Document
        path = export(_resume(), str(tmp_path / "r"), "docx")
        doc = Document(path)
        bolded = any(run.bold for p in doc.paragraphs for run in p.runs)
        assert bolded


class TestStore:
    def test_crud_default_and_versions(self, dbpath):
        from core.resume import (default_resume, get_resume, list_resumes,
                                 save_resume, save_version, set_default)
        r = _resume()
        rid = save_resume("Base", "raw text " * 50, r)
        assert get_resume(rid)["is_default"] is True   # first becomes default
        rid2 = save_resume("Second", "raw " * 50, r)
        set_default(rid2)
        assert default_resume()["id"] == rid2
        vid = save_version(rid, r, label="Tailored for job 1")
        from core.resume.store import list_versions
        assert any(v["id"] == vid for v in list_versions(rid))
